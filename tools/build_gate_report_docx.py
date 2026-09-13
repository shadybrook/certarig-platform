#!/usr/bin/env python3
"""Create the polished DOCX version of the CertaRig gate evidence review."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "explainer" / "2026-09-13-gates-0-to-5-analysis.md"
OUT = ROOT / "deliverables" / "CertaRig_Gates_0_to_5_Evidence_and_Product_Direction.docx"

BLACK = "000000"
INK = "172523"
MUTED = "5F6B68"
DARK_GREEN = "1F5E4C"
PALE_GREEN = "E8F2ED"
PALE_BLUE = "EAF1F5"
GRID = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 110, start: int = 120, bottom: int = 110, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_font(run, name: str = "Arial", size: float | None = None, color: str | None = None,
             bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_inline(paragraph, content: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(content):
        if match.start() > pos:
            set_font(paragraph.add_run(content[pos:match.start()]), size=11)
        token = match.group(0)
        if token.startswith("**"):
            set_font(paragraph.add_run(token[2:-2]), size=11, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Courier New", size=9.5, color=INK)
        pos = match.end()
    if pos < len(content):
        set_font(paragraph.add_run(content[pos:]), size=11)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_font(run, size=9, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, before, after in (
        ("Title", 26, 0, 10),
        ("Heading 1", 19, 16, 8),
        ("Heading 2", 14, 12, 6),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("CERTARIG  ·  GATE EVIDENCE REVIEW"), size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("13 September 2026   ·   "), size=9, color=MUTED)
    add_page_number(footer)


def cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CertaRig crossed the lab boundary\nwithout giving the agent control of safety")
    set_font(r, size=25, bold=True, color=BLACK)
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Gate 0 to Gate 5 evidence review and product direction")
    set_font(r, size=14, color=DARK_GREEN, bold=True)
    p.paragraph_format.space_after = Pt(30)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Six numbered gates  ·  Six passing hardware runs  ·  3,422 recorded samples")
    set_font(r, size=11.5, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Prepared from the CertaRig_CursorControl repository\n13 September 2026")
    set_font(r, size=10.5, color=MUTED)
    p.paragraph_format.space_before = Pt(22)

    doc.add_page_break()


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) >= 2 and all(set(cell) <= {"-", ":", " "} for cell in rows[1]):
        rows.pop(1)
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.42), Inches(3.2), Inches(1.85)] if len(rows[0]) == 3 else [Inches(6.7 / len(rows[0]))] * len(rows[0])
    for col, value in enumerate(rows[0]):
        cell = table.rows[0].cells[col]
        cell.width = widths[col]
        set_cell_fill(cell, DARK_GREEN)
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_font(p.add_run(value), size=9.5, color=WHITE, bold=True)
    for row_idx, values in enumerate(rows[1:]):
        cells = table.add_row().cells
        for col, value in enumerate(values):
            cell = cells[col]
            cell.width = widths[col]
            set_cell_fill(cell, PALE_GREEN if row_idx % 2 else WHITE)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            add_inline(p, value)
            for run in p.runs:
                if run.font.size is None or run.font.size.pt > 9.5:
                    run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc: Document, alt: str, rel_path: str, source: Path = SOURCE) -> None:
    path = (source.parent / rel_path).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.65))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    set_font(caption.add_run(alt), size=9, color=MUTED, italic=True)


def add_code(doc: Document, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_fill(cell, PALE_BLUE)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    set_cell_borders(cell, color=GRID)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("\n".join(lines)), name="Courier New", size=9.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def body_from_markdown(doc: Document, source: Path = SOURCE) -> None:
    lines = source.read_text().splitlines()
    idx = 1
    skip_metadata = True
    break_before = {
        "What each gate established",
        "What the data says",
        "The product pivot",
        "Evidence limitations and corrective actions",
        "Reproducibility and provenance",
    }
    while idx < len(lines):
        line = lines[idx]
        stripped = line.strip()
        if skip_metadata and (not stripped or stripped.startswith("**Evidence review") or stripped.startswith("Prepared ")):
            idx += 1
            continue
        skip_metadata = False
        if not stripped:
            idx += 1
            continue
        if stripped.startswith("## "):
            heading = stripped[3:]
            if heading in break_before:
                doc.add_page_break()
            doc.add_heading(heading, level=1)
            idx += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            idx += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", stripped)
        if image_match:
            add_image(doc, image_match.group(1), image_match.group(2), source)
            idx += 1
            continue
        if stripped.startswith("```"):
            code_lines: list[str] = []
            idx += 1
            while idx < len(lines) and not lines[idx].strip().startswith("```"):
                code_lines.append(lines[idx])
                idx += 1
            idx += 1
            add_code(doc, code_lines)
            continue
        if stripped.startswith("|"):
            table_lines: list[str] = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                table_lines.append(lines[idx])
                idx += 1
            add_table(doc, table_lines)
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.right_indent = Inches(0.5)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(10)
            quote = stripped[2:]
            if quote.startswith("**") and quote.endswith("**"):
                quote = quote[2:-2]
            run = p.add_run(quote)
            set_font(run, size=13, color=DARK_GREEN, bold=True, italic=True)
            idx += 1
            continue
        if re.match(r"^\d+\. ", stripped):
            match = re.match(r"^(\d+)\. (.*)", stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            set_font(p.add_run(f"{match.group(1)}.  "), size=11, bold=True, color=DARK_GREEN)
            add_inline(p, match.group(2))
            idx += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, stripped[2:])
            idx += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, stripped)
        idx += 1


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    cover(doc)
    body_from_markdown(doc)
    core = doc.core_properties
    core.title = "CertaRig Gates 0 to 5 Evidence and Product Direction"
    core.subject = "Evidence review of the 12 September 2026 CertaRig dry bench lab"
    core.author = "Chintan Dedhia"
    core.keywords = "CertaRig, commissioning, evidence, guardrail, Raspberry Pi"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
