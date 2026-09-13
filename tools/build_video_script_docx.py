#!/usr/bin/env python3
"""Create a reader friendly DOCX of the CertaRig visual explainer script."""

from pathlib import Path

from build_gate_report_docx import body_from_markdown, configure_document, set_font
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "explainer" / "2026-09-13-video-script.md"
OUT = ROOT / "deliverables" / "CertaRig_Visual_Explainer_Video_Script.docx"


def cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("THE AI CAN ASK.\nONLY THE KERNEL CAN SAY YES.")
    set_font(r, size=27, bold=True, color="000000")
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("CertaRig visual explainer script"), size=15, color="1F5E4C", bold=True)
    p.paragraph_format.space_after = Pt(26)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("Target duration 10:30 to 12:00\nOriginal curiosity led engineering narrative"), size=11, color="5F6B68")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    set_font(p.add_run("Prepared 13 September 2026"), size=10.5, color="5F6B68")
    doc.add_page_break()


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    header = doc.sections[0].header.paragraphs[0]
    for run in header.runs:
        run.text = ""
    set_font(header.add_run("CERTARIG  ·  VISUAL EXPLAINER SCRIPT"), size=8.5, color="5F6B68", bold=True)
    cover(doc)
    body_from_markdown(doc, SOURCE)
    doc.core_properties.title = "CertaRig Visual Explainer Script"
    doc.core_properties.author = "Chintan Dedhia"
    doc.core_properties.subject = "CertaRig gate evidence and product direction explainer"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
