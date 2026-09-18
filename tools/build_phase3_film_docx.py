#!/usr/bin/env python3
"""Create a reader-friendly DOCX of the Phase 3 submission film script."""

from pathlib import Path

from build_gate_report_docx import body_from_markdown, configure_document, set_font
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "explainer" / "phase3-film" / "SCRIPT_EDIT_ME.md"
OUT = ROOT / "deliverables" / "CertaRig_Phase3_Submission_Film_Script.docx"


def cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WHAT IS CERTARIG?")
    set_font(r, size=28, bold=True, color="000000")
    p.paragraph_format.space_after = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("The AI can ask. Only the kernel can say yes."), size=14, color="1F5E4C", bold=True)
    p.paragraph_format.space_after = Pt(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(
        p.add_run("Phase 3 submission film script\nTarget duration 11:40 to 12:10\nCurastra-grammar product film"),
        size=11,
        color="5F6B68",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    set_font(p.add_run("Prepared 16 September 2026"), size=10.5, color="5F6B68")
    doc.add_page_break()


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    header = doc.sections[0].header.paragraphs[0]
    for run in header.runs:
        run.text = ""
    set_font(header.add_run("CERTARIG  ·  PHASE 3 SUBMISSION FILM"), size=8.5, color="5F6B68", bold=True)
    footer = doc.sections[0].footer.paragraphs[0]
    for run in footer.runs:
        if run.text.startswith("13 September"):
            run.text = "16 September 2026   ·   "
    cover(doc)
    body_from_markdown(doc, SOURCE)
    doc.core_properties.title = "CertaRig Phase 3 Submission Film Script"
    doc.core_properties.author = "Chintan Dedhia"
    doc.core_properties.subject = "Timed voiceover, stills, and demo stitch plan"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
